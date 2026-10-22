"""Messy real-world Excel layouts: stacked tables, numbers stored as text, months down the rows, units notes."""
import datetime as dt

import pytest
from openpyxl import Workbook

from common.units import detect_currency
from ingest_financials import ingest_excel, parse_number_text

MONTHS = [dt.date(2026, 1, 1), dt.date(2026, 2, 1), dt.date(2026, 3, 1)]


def _messy(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MIS"
    ws["A1"] = "Kestrel Foods Pvt Ltd - Management Accounts"
    ws.merge_cells("A1:F1")
    ws["A2"] = "All figures in ₹ lakhs"
    ws.append([])
    ws.append(["Particulars", None, *MONTHS, "Q4 FY26 Total"])
    ws.append(["Revenue", None, None, None, None, None])
    ws.append(["Gross Sales", None, "1,23,456.50", "1,30,210.00", "1,41,005.25", "3,94,671.75"])
    ws.append(["Less: Discounts", None, "(12,345.00)", "(13,020.10)", "(14,100.50)", "(39,465.60)"])
    ws.append([])
    ws.append([])
    ws.append(["Customer KPIs"])
    ws.append(["Metric", None, *MONTHS, None])
    ws.append(["New Customers", None, 1200, 1350, 1500, None])
    ws2 = wb.create_sheet("Cash (monthly)")
    ws2.append(["Cash Summary (INR)"])
    ws2.append(["Month", "Cash In", "Cash Out", "Closing Cash"])
    for m, ci, co, cc in zip(MONTHS, [5_200_000, 5_350_000, 5_600_000], [6_100_000, 6_000_000, 6_250_000],
                             [48_000_000, 47_350_000, 46_700_000]):
        ws2.append([m, ci, co, cc])
    wb.save(path)
    return path


@pytest.mark.parametrize("text,value", [("1,23,456.50", 123456.5), ("(12,345.00)", -12345.0), ("₹ 1,200", 1200.0),
                                        ("-450", -450.0), ("Rs. 2,50,000", 250000.0), ("(1,000", None),
                                        ("Q4 FY26", None), ("12,34", None), ("n/a", None)])
def test_parse_number_text(text, value):
    assert parse_number_text(text) == value


def test_detect_currency():
    assert detect_currency("Cash Summary (INR)") == "INR"
    assert detect_currency("Amounts in USD") == "USD"
    assert detect_currency("Currency: INR") == "INR"
    assert detect_currency("Revenue grew in March") is None


def test_messy_workbook(tmp_path):
    chunks = ingest_excel(_messy(tmp_path / "messy.xlsx"))
    by_section = {c.metadata.section: c for c in chunks}
    rev = by_section["Revenue"]
    assert rev.metadata.units == "INR lakh" and rev.metadata.periods == ["Jan-26", "Feb-26", "Mar-26"]
    assert "| Gross Sales | 123,456.50 | 130,210.00 | 141,005.25 | 394,671.75 |" in rev.text
    assert "| Less: Discounts | -12,345.00 |" in rev.text
    kpi = by_section["Customer KPIs"]           # second table on the sheet, own header row
    assert "2026-01-01" not in kpi.text and "| New Customers | 1,200 | 1,350 | 1,500 |" in kpi.text
    assert kpi.metadata.units == ""             # KPI table does not inherit "All figures in ₹ lakhs"
    assert "Col" not in kpi.text
    cash = by_section["Cash (monthly)"]         # months down the rows -> transposed
    assert cash.metadata.periods == ["Jan-26", "Feb-26", "Mar-26"] and cash.metadata.units == "INR"
    assert "| Cash Out | 6,100,000 | 6,000,000 | 6,250,000 |" in cash.text
    assert len({c.chunk_id for c in chunks}) == len(chunks)
