"""Table-aware ingestion: Excel (openpyxl + pandas) and PDF (Docling) -> Markdown table chunks.

Rules (from CLAUDE.md):
- Every chunk carries its column headers, row labels, period(s), units and source metadata.
- A value is never separated from its header: long sections are split by rows and each piece repeats
  the header row.
- Chunks are rendered with common.tables.render_chunk, the same format the model was trained on.

Excel layout expected (typical MIS export):
  optional title rows ("Monthly P&L", "Units: INR")
  header row: first cell = row-label header ("Line Item"), other cells = periods or column names
  data rows; a row with a label but no values starts a new section ("Revenue", "Variable Costs")

Usage:
  python ingest_financials.py sample_startup_mis.xlsx [more.xlsx board_pack.pdf ...] [--out chunks.jsonl]
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from collections.abc import Iterable, Sequence
from pathlib import Path

import pandas as pd
import yaml

from common.schemas import ChunkMetadata, TableChunk
from common.tables import format_cell, render_chunk, render_text_chunk
from common.units import detect_units, first_currency, first_units

# Body text only sets units when it reads like a units statement ("All figures in INR lakh unless stated").
UNITS_CUE_RE = re.compile(r"figures|amounts|values|unless|units|currency|\(\s*(?:₹|rs\.?|inr|usd|\$)", re.IGNORECASE)

ROOT = Path(__file__).resolve().parent

PERIOD_RE = re.compile(
    r"^(?:[A-Z][a-z]{2}-\d{2}|[A-Z][a-z]{2} \d{4}|FY\s?\d{2,4}(?:[- ]?Q[1-4])?|Q[1-4][- ]?(?:FY)?\s?\d{2,4}"
    r"|\d{4}-\d{2}|H[12][- ]?(?:FY)?\d{2,4}|FY\d{2} Plan|Base|Bull|Bear)$"
)
UNITS_RE = re.compile(r"(?:units?|currency)\s*[:=]\s*(.+)$", re.IGNORECASE)


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:40]


def _header_text(v: object) -> str:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%b-%y")
    return format_cell(v)


def _is_empty(v: object) -> bool:
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and not v.strip())


def split_sections(rows: Sequence[Sequence[str]], default: str) -> list[tuple[str, list[list[str]]]]:
    """Group data rows into sections. A row with a label and no values is a section header."""
    sections: list[tuple[str, list[list[str]]]] = []
    name, cur = default, []
    for r in rows:
        label, values = r[0], r[1:]
        if label and all(not v for v in values):
            if cur:
                sections.append((name, cur))
            name, cur = label, []
            continue
        if not label and all(not v for v in values):
            continue
        cur.append(list(r))
    if cur:
        sections.append((name, cur))
    return sections


def make_chunks(*, title: str, source: str, sheet: str, headers: list[str], rows: list[list[str]],
                units: str, default_section: str, extra_periods: Iterable[str] = (), page: int | None = None,
                max_rows: int = 12) -> list[TableChunk]:
    periods = [h for h in headers[1:] if PERIOD_RE.match(h)] or list(extra_periods)
    chunks: list[TableChunk] = []
    for section, sec_rows in split_sections(rows, default_section):
        for start in range(0, len(sec_rows), max_rows):
            piece = sec_rows[start:start + max_rows]
            ctitle = f"{title} - {section}" if section != title else title
            meta = ChunkMetadata(source=source, sheet=sheet, section=section, periods=periods,
                                 row_labels=[r[0] for r in piece], units=units, page=page)
            cid = f"{_slug(Path(source).stem)}::{_slug(sheet)}::{_slug(section)}::{start // max_rows}"
            chunks.append(TableChunk(chunk_id=cid, title=ctitle, text=render_chunk(ctitle, meta, headers, piece),
                                     metadata=meta))
    return chunks


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _is_period(v: object) -> bool:
    return isinstance(v, (dt.datetime, dt.date)) or (isinstance(v, str) and bool(PERIOD_RE.match(v.strip())))


# Numbers stored as text: "1,23,456.50" (Indian grouping), "(12,345)" (negative), "₹ 1,200", "-450".
_NUM_TEXT_RE = re.compile(r"^(\()?\s*(?:₹|rs\.?|inr|\$|usd)?\s*(-)?\s*(\d{1,3}(?:,\d{3})+|\d{1,2}(?:,\d{2})*,\d{3}|\d+)(\.\d+)?\s*(\))?$",
                          re.IGNORECASE)


def parse_number_text(s: str) -> float | None:
    """'1,23,456.50' -> 123456.5; '(12,345.00)' -> -12345.0; anything else -> None."""
    m = _NUM_TEXT_RE.match(s.strip())
    if not m or bool(m.group(1)) != bool(m.group(5)):
        return None
    v = float(m.group(3).replace(",", "") + (m.group(4) or ""))
    return -v if (m.group(1) or m.group(2)) else v


def _cell_text(v: object, is_pct: bool) -> str:
    """A data cell as the chunk shows it: dates as Mon-YY, numbers-as-text normalised to plain numbers."""
    if _is_empty(v):
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%b-%y")
    if isinstance(v, str):
        n = parse_number_text(v)
        return format_cell(n) if n is not None else format_cell(v)
    return format_cell(v, is_pct)


# A block whose title reads like operational KPIs does not inherit a sheet-wide "All figures in INR lakh" note.
NON_MONEY_TITLE_RE = re.compile(r"kpi|metric|customer|headcount|operational|funnel|users", re.IGNORECASE)


def _header_row(values: pd.DataFrame) -> object | None:
    """Fallback header: first row whose first cell is text and that has >= 2 non-empty cells."""
    for i, row in values.iterrows():
        if isinstance(row.iloc[0], str) and sum(not _is_empty(v) for v in row) >= 2:
            return i
    return None


def _excel_blocks(values: pd.DataFrame) -> list[tuple[list, object, list]]:
    """Split one sheet into tables: (title row indexes, header row index, data row indexes).

    A header is a row with >= 2 period cells (month dates, "Mar-26", "Q4 FY26", ...). Several headers = several
    stacked tables; the label-only rows right above a later header are that table's title rows. Without any
    period header the whole sheet is one table with the fallback header (Pricing-style "Item | Value" sheets,
    or months-down-the-rows tables that get transposed later)."""
    idx = list(values.index)
    headers = [i for i in idx if sum(_is_period(v) for v in values.loc[i].iloc[1:]) >= 2]
    if not headers:
        h = _header_row(values)
        return [] if h is None else [([i for i in idx if i < h], h, [i for i in idx if i > h])]
    pos = {i: k for k, i in enumerate(idx)}
    titles: list[list] = []
    for n, h in enumerate(headers):
        if n == 0:
            titles.append([i for i in idx if i < h])
            continue
        t: list = []
        p = pos[h] - 1
        while p > pos[headers[n - 1]] and sum(not _is_empty(v) for v in values.loc[idx[p]]) <= 1:
            t.insert(0, idx[p])
            if p == 0 or idx[p] - idx[p - 1] > 1:  # a blank line above: the title block starts here
                break
            p -= 1
        titles.append(t)
    blocks = []
    for n, h in enumerate(headers):
        end = titles[n + 1][0] if n + 1 < len(headers) and titles[n + 1] else (headers[n + 1] if n + 1 < len(headers) else None)
        data = [i for i in idx if i > h and (end is None or i < end)]
        blocks.append((titles[n], h, data))
    return blocks


def ingest_excel(path: Path, max_rows: int = 12) -> list[TableChunk]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)  # cached values, not formulas
    chunks: list[TableChunk] = []
    for ws in wb.worksheets:
        # pandas frame of raw values (structure), parallel frame of "is percent" flags (formatting)
        values = pd.DataFrame([[c.value for c in row] for row in ws.iter_rows()])
        if values.empty:
            continue
        pct = pd.DataFrame([["%" in (c.number_format or "") for c in row] for row in ws.iter_rows()])
        keep_cols = [c for c in values.columns if not values[c].map(_is_empty).all()]
        values, pct = values[keep_cols], pct[keep_cols]
        values = values[~values.map(_is_empty).all(axis=1)]
        pct = pct.loc[values.index]
        if values.empty:
            continue

        sheet_units = ""
        for n, (title_idx, header_idx, data_idx) in enumerate(_excel_blocks(values)):
            title_rows = [str(values.loc[i].iloc[0]).strip() for i in title_idx if not _is_empty(values.loc[i].iloc[0])]
            title, units, extra_periods = ws.title, "", []
            for t in title_rows:
                m = UNITS_RE.search(t)
                if m:
                    units = m.group(1).strip()
                elif PERIOD_RE.match(t):
                    extra_periods.append(t)
                elif title == ws.title:
                    title = t
            header_vals = list(values.loc[header_idx])
            headers = [_header_text(v) or f"Col{j}" for j, v in enumerate(header_vals)]
            rows = [[_cell_text(v, bool(pct.loc[i, c])) if j else (_header_text(v) if not _is_empty(v) else "")
                     for j, (c, v) in enumerate(values.loc[i].items())] for i in data_idx]
            # a column that only another table on the sheet uses: no header and empty in this table
            cols = [j for j in range(len(headers)) if j == 0 or not _is_empty(header_vals[j]) or any(r[j] for r in rows)]
            headers, rows = [headers[j] for j in cols], [[r[j] for j in cols] for r in rows]
            # Months down the rows (Month | Cash In | Cash Out): transpose to Line Item | Jan-26 | Feb-26 | ...
            firsts = [values.loc[i].iloc[0] for i in data_idx]
            if (sum(_is_period(v) for v in headers[1:]) < 2 and firsts
                    and sum(_is_period(v) for v in firsts) >= max(2, 0.6 * len(firsts))):
                keep = [r for r, v in zip(rows, firsts) if _is_period(v)]
                headers, rows = (["Line Item", *[r[0] for r in keep]],
                                 [[headers[j], *[r[j] for r in keep]] for j in range(1, len(headers))])
            # Scaled units: "Units: INR Lakhs" -> "INR lakh"; else a title row / header cell / sheet name stating
            # them ("P&L (₹ in lakhs)", "Particulars (INR Cr)"); else a plain currency in the title ("(INR)").
            units = detect_units(units) or first_units([*title_rows, *headers, ws.title]) or units
            if n == 0:
                sheet_units = units
            elif not units and not NON_MONEY_TITLE_RE.search(" ".join(title_rows)):
                units = sheet_units  # a sheet-wide units note covers later tables too
            units = units or first_currency(title_rows) or ""
            chunks += make_chunks(title=title, source=path.name, sheet=ws.title, headers=headers, rows=rows, units=units,
                                  default_section=ws.title if n == 0 else title, extra_periods=extra_periods,
                                  max_rows=max_rows)
    seen: dict[str, int] = {}
    for c in chunks:  # two tables on one sheet may share a section name
        k = seen.get(c.chunk_id, 0)
        seen[c.chunk_id] = k + 1
        if k:
            c.chunk_id = f"{c.chunk_id}::{k}"
    return chunks


# ---------------------------------------------------------------------------
# PDF (Docling)
# ---------------------------------------------------------------------------

def ingest_pdf(path: Path, max_rows: int = 12, text_chunk_chars: int = 1200) -> list[TableChunk]:
    from docling.document_converter import DocumentConverter
    from docling_core.types.doc import SectionHeaderItem, TableItem, TextItem

    doc = DocumentConverter().convert(str(path)).document
    chunks: list[TableChunk] = []
    section, buf, buf_page = "Document", [], None
    recent: list[str] = []          # last text items of the current section (a units note sits just above a table)
    doc_units: str | None = None    # "All figures in INR lakh" anywhere earlier applies to later tables

    def flush_text() -> None:
        nonlocal buf
        text = "\n".join(buf).strip()
        if text:
            meta = ChunkMetadata(source=path.name, sheet="text", section=section, page=buf_page)
            cid = f"{_slug(path.stem)}::text::{_slug(section)}::{len(chunks)}"
            chunks.append(TableChunk(chunk_id=cid, title=section, text=render_text_chunk(section, meta, text), metadata=meta))
        buf = []

    for item, _level in doc.iterate_items():
        page = item.prov[0].page_no if getattr(item, "prov", None) else None
        if isinstance(item, TableItem):
            df = item.export_to_dataframe(doc=doc)
            if df.empty:
                continue
            headers = [_header_text(c) for c in df.columns]
            rows = [[format_cell(v) for v in r] for r in df.itertuples(index=False)]
            caption = item.caption_text(doc) or section
            units = (first_units([item.caption_text(doc), *headers, section])
                     or first_units(t for t in reversed(recent) if UNITS_CUE_RE.search(t)) or doc_units or "")
            chunks += make_chunks(title=caption, source=path.name, sheet=f"page {page}", headers=headers, rows=rows,
                                  units=units, default_section=caption, page=page, max_rows=max_rows)
        elif isinstance(item, SectionHeaderItem):
            flush_text()
            section = item.text.strip() or section
            recent = []
        elif isinstance(item, TextItem):
            if not buf:
                buf_page = page
            buf.append(item.text)
            recent = (recent + [item.text])[-3:]
            if UNITS_CUE_RE.search(item.text):
                doc_units = detect_units(item.text) or doc_units
            if sum(len(b) for b in buf) > text_chunk_chars:
                flush_text()
    flush_text()
    return chunks


# ---------------------------------------------------------------------------

def ingest_paths(paths: Iterable[str | Path], max_rows: int = 12) -> list[TableChunk]:
    out: list[TableChunk] = []
    for p in map(Path, paths):
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            out += ingest_excel(p, max_rows)
        elif suffix == ".pdf":
            out += ingest_pdf(p, max_rows)
        else:
            raise ValueError(f"unsupported file type: {p}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--config", default=str(ROOT / "configs" / "rag.yaml"))
    ap.add_argument("--out", help="write chunks as JSONL (keep real data inside git-ignored folders)")
    args = ap.parse_args()
    cfg = yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))
    chunks = ingest_paths(args.paths, cfg["max_rows_per_chunk"])
    for c in chunks:
        print(f"\n=== {c.chunk_id}\n{c.text}")
    print(f"\n{len(chunks)} chunks", file=sys.stderr)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            for c in chunks:
                f.write(c.model_dump_json() + "\n")


if __name__ == "__main__":
    main()
